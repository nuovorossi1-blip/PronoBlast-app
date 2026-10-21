"""Quote della scheda partita Sisal (21/10/2026): multigol, combo, Under 4,5, risultati esatti.

Il PDF del palinsesto ha solo i mercati principali: il resto sta nella pagina della partita
su sisal.it. Qui si trova il codice Sisal (colonne Pal. e Avv.) negli Excel gia' scaricati,
poi Edge legge la scheda dal servizio che usa il sito.
Sisal blocca Edge pilotato da Playwright (Access Denied): si apre un Edge NORMALE con un
profilo tutto suo e ci si collega dopo (CDP). Una sola apertura per tutte le partite chieste.

Prova:  python scheda.py "2026-10-24" "Napoli" "Roma"
"""
import datetime, json, pathlib, re, shutil, socket, subprocess, sys, time

from openpyxl import load_workbook

PORTA = 9335
HOME = "https://www.sisal.it/scommesse-matchpoint"
API = "https://betting.sisal.it/api/lettura-palinsesto-sport/palinsesto/prematch/schedaAvvenimento/"
# Mercati dei giocatori: centinaia di righe che il motore non usa.
SALTA = ("MARCATORE", "SEGNA O SUO SOSTITUTO", "GIOCATORE", "ASSIST", "TIRI ", "CARTELLIN", "AMMONIT")
MESI = {m: i + 1 for i, m in enumerate("gennaio febbraio marzo aprile maggio giugno luglio agosto settembre ottobre novembre dicembre".split())}


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def codici_da_excel(cartella: pathlib.Path, partite: list[dict], quanti=6) -> dict:
    """{id: 'pal-avv'} cercando giorno + squadre negli ultimi Excel (dal piu' recente)."""
    cerca = {(p["day"], _norm(p["squadra1"]), _norm(p["squadra2"])): p["id"] for p in partite}
    trovati = {}
    for f in sorted(cartella.glob("*.xlsx"), key=lambda x: x.stat().st_mtime, reverse=True)[:quanti]:
        anno = datetime.date.today().year
        try:
            wb = load_workbook(f, read_only=True, data_only=True)
        except Exception:
            continue
        for ws in wb.worksheets:
            giorno, ultimo_mese = None, None
            for r in ws.iter_rows(values_only=True):
                r = list(r) + [None] * 8
                d = re.search(r"(\d{1,2})\s+(" + "|".join(MESI) + r")", str(r[0] or ""), re.I)
                if d:
                    mese = MESI[d[2].lower()]
                    if ultimo_mese and ultimo_mese - mese > 6:
                        anno += 1
                    ultimo_mese = mese
                    giorno = f"{anno}-{mese:02d}-{int(d[1]):02d}"
                    continue
                if giorno and r[2] and r[3]:
                    k = (giorno, _norm(r[4]), _norm(r[5]))
                    if k in cerca and cerca[k] not in trovati:
                        trovati[cerca[k]] = f"{int(r[2])}-{int(r[3])}"
        wb.close()
        if len(trovati) == len(cerca):
            break
    return trovati


def _edge():
    for p in (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe", r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if pathlib.Path(p).exists():
            return p
    return shutil.which("msedge") or "msedge"


def _attendi_porta(sec=20):
    fine = time.time() + sec
    while time.time() < fine:
        try:
            with socket.create_connection(("127.0.0.1", PORTA), 0.5):
                return True
        except OSError:
            time.sleep(0.3)
    return False


def _normalizza(d: dict) -> list:
    nomi = {v["codiceScommessa"]: (v.get("descrizione") or "").strip() for v in d.get("scommessaMap", {}).values()}
    mercati = []
    for v in d.get("infoAggiuntivaMap", {}).values():
        nome, desc = nomi.get(v["codiceScommessa"], ""), (v.get("descrizione") or "").strip()
        if any(s in (nome + " " + desc).upper() for s in SALTA):
            continue
        esiti = [[(e.get("descrizione") or "").strip(), round(e["quota"] / 100, 2)]
                 for e in v.get("esitoList", []) if e.get("quota") and e.get("stato", 1) == 1]
        if esiti:
            mercati.append({"cs": v["codiceScommessa"], "nome": nome, "desc": desc, "soglia": (v.get("soglia") or "").strip(), "esiti": esiti})
    return mercati


def leggi_schede(codici: list[str], profilo: pathlib.Path) -> dict:
    """{codice: {ok, at, evento, mercati}}"""
    from playwright.sync_api import sync_playwright
    profilo.mkdir(parents=True, exist_ok=True)
    proc = subprocess.Popen([_edge(), f"--remote-debugging-port={PORTA}", f"--user-data-dir={profilo}", "--no-first-run",
                             "--no-default-browser-check", "--window-position=-32000,-32000", "--window-size=1280,900", "about:blank"])
    out = {}
    try:
        if not _attendi_porta():
            raise RuntimeError("Edge non si e' aperto")
        with sync_playwright() as p:
            b = p.chromium.connect_over_cdp(f"http://127.0.0.1:{PORTA}")
            ctx = b.contexts[0]
            pg = ctx.pages[0] if ctx.pages else ctx.new_page()
            pg.goto(HOME, timeout=90_000)
            pg.wait_for_timeout(3000)
            if "denied" in (pg.title() or "").lower():
                raise RuntimeError("Sisal ha bloccato Edge (Access Denied)")
            for c in codici:
                r = pg.evaluate("""async u=>{const r=await fetch(u,{credentials:'include'});return {s:r.status,t:await r.text()}}""", API + c)
                if r["s"] != 200:
                    out[c] = {"ok": False, "errore": f"Sisal HTTP {r['s']}"}
                    continue
                try:
                    d = json.loads(r["t"])
                    out[c] = {"ok": True, "evento": (d.get("avvenimentoFe") or {}).get("descrizione"), "mercati": _normalizza(d)}
                except Exception as e:
                    out[c] = {"ok": False, "errore": f"risposta non leggibile: {e}"}
                out[c]["at"] = datetime.datetime.now(datetime.timezone.utc).isoformat()
            b.close()
    finally:
        proc.terminate()
        try:
            proc.wait(10)
        except Exception:
            proc.kill()
    return out


def schede_partite(partite: list[dict], cartella_excel: pathlib.Path, profilo: pathlib.Path) -> dict:
    """{id: {ok, codice, at, evento, mercati} | {ok: False, errore}}"""
    codici = codici_da_excel(cartella_excel, partite)
    esito = {p["id"]: {"ok": False, "errore": "partita non trovata negli Excel Sisal scaricati"} for p in partite if p["id"] not in codici}
    if codici:
        lette = leggi_schede(sorted(set(codici.values())), profilo)
        for pid, c in codici.items():
            esito[pid] = {**lette.get(c, {"ok": False, "errore": "scheda non letta"}), "codice": c}
    return esito


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    import agente
    giorno, casa, ospite = sys.argv[1:4]
    r = schede_partite([{"id": "prova", "day": giorno, "squadra1": casa, "squadra2": ospite}], agente.CARTELLA_EXCEL, agente.BASE / "edge-sisal")
    v = r["prova"]
    print(json.dumps({k: (len(x) if k == "mercati" else x) for k, x in v.items()}, ensure_ascii=False))
